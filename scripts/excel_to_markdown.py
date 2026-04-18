#!/usr/bin/env python3
"""
Excel to Markdown Converter

Converts Excel (.xlsx) and CSV files in a given directory to readable
Markdown format. Used by Claude Code agents to process tabular input data.

Usage:
    python3 scripts/excel_to_markdown.py <input_directory>

Output:
    Prints Markdown-formatted tables to stdout for each file found.
"""

import csv
import os
import sys


def read_csv_file(filepath):
    """Read a CSV file and return rows."""
    rows = []
    encodings = ["utf-8", "gbk", "gb2312", "utf-8-sig", "latin-1"]
    for enc in encodings:
        try:
            with open(filepath, "r", encoding=enc) as f:
                reader = csv.reader(f)
                rows = [row for row in reader]
            break
        except (UnicodeDecodeError, UnicodeError):
            continue
    return rows


def read_excel_file(filepath):
    """Read an Excel file and return a dict of sheet_name -> rows."""
    try:
        import openpyxl
    except ImportError:
        print(f"⚠️ openpyxl not installed. Install with: pip install openpyxl", file=sys.stderr)
        print(f"Attempting to install openpyxl...", file=sys.stderr)
        os.system("pip install openpyxl")
        import openpyxl

    sheets = {}
    wb = openpyxl.load_workbook(filepath, read_only=True, data_only=True)
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        rows = []
        for row in ws.iter_rows(values_only=True):
            str_row = [str(cell) if cell is not None else "" for cell in row]
            if any(cell.strip() for cell in str_row):
                rows.append(str_row)
        if rows:
            sheets[sheet_name] = rows
    wb.close()
    return sheets


def rows_to_markdown(rows, title=None):
    """Convert a list of rows to a Markdown table."""
    if not rows:
        return ""

    lines = []
    if title:
        lines.append(f"### {title}")
        lines.append("")

    num_cols = max(len(row) for row in rows)
    normalized = []
    for row in rows:
        padded = list(row) + [""] * (num_cols - len(row))
        normalized.append(padded)

    col_widths = [0] * num_cols
    for row in normalized:
        for i, cell in enumerate(row):
            col_widths[i] = max(col_widths[i], len(cell))

    def format_row(row):
        cells = [cell.ljust(col_widths[i]) for i, cell in enumerate(row)]
        return "| " + " | ".join(cells) + " |"

    lines.append(format_row(normalized[0]))
    lines.append("|" + "|".join("-" * (w + 2) for w in col_widths) + "|")

    for row in normalized[1:]:
        lines.append(format_row(row))

    lines.append("")
    return "\n".join(lines)


def process_directory(input_dir):
    """Process all Excel and CSV files in the given directory."""
    if not os.path.isdir(input_dir):
        print(f"Error: Directory '{input_dir}' not found.", file=sys.stderr)
        sys.exit(1)

    files = sorted(os.listdir(input_dir))
    found = False

    for filename in files:
        filepath = os.path.join(input_dir, filename)

        if not os.path.isfile(filepath):
            continue

        ext = os.path.splitext(filename)[1].lower()

        if ext == ".csv":
            found = True
            print(f"## 文件: {filename}\n")
            rows = read_csv_file(filepath)
            if rows:
                print(rows_to_markdown(rows))
            else:
                print(f"（文件为空或无法解析）\n")

        elif ext in (".xlsx", ".xls"):
            found = True
            print(f"## 文件: {filename}\n")
            try:
                sheets = read_excel_file(filepath)
                for sheet_name, rows in sheets.items():
                    print(rows_to_markdown(rows, title=f"Sheet: {sheet_name}"))
            except Exception as e:
                print(f"（解析失败: {e}）\n", file=sys.stderr)

    if not found:
        print("（目录中未找到 Excel 或 CSV 文件）")


def main():
    if len(sys.argv) < 2:
        print(f"Usage: {sys.argv[0]} <input_directory>", file=sys.stderr)
        sys.exit(1)

    input_dir = sys.argv[1]
    process_directory(input_dir)


if __name__ == "__main__":
    main()
