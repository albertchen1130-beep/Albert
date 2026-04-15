"""
Agent 3: ETF数据分析 (Tables 4-5)
Reads ETF Excel data, extracts key metrics, and generates
Tables 4 and 5 for the biweekly report.
"""

import os
from scripts.agents.base import BaseAgent


def parse_excel_to_text(excel_path: str) -> str:
    """Parse an Excel file into text representation for the LLM."""
    try:
        import openpyxl
    except ImportError:
        return "[ERROR] openpyxl not installed. Run: pip install openpyxl"

    if not os.path.exists(excel_path):
        return f"[ERROR] Excel file not found: {excel_path}"

    wb = openpyxl.load_workbook(excel_path, data_only=True)
    sections = []

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        lines = [f"## Sheet: {sheet_name}"]

        for row in ws.iter_rows(min_row=1, values_only=False):
            cells = []
            for cell in row:
                val = cell.value
                if val is None:
                    cells.append("")
                elif isinstance(val, float):
                    # Format percentages and numbers nicely
                    if abs(val) < 1 and val != 0:
                        cells.append(f"{val:.4%}")
                    else:
                        cells.append(f"{val:,.2f}")
                else:
                    cells.append(str(val))
            lines.append(" | ".join(cells))

        sections.append("\n".join(lines))

    wb.close()
    return "\n\n".join(sections)


class ETFAgent(BaseAgent):
    AGENT_NAME = "ETFAgent (表4-5 ETF分析)"
    SYSTEM_PROMPT = """你是一位资深的ETF产品分析师，专门负责编写双周报中"表4"和"表5"的ETF数据分析板块。

你的任务：
1. 接收从Excel中提取的ETF数据
2. 分析数据并生成表4和表5的规范内容
3. 包含关键指标的变动分析

表4通常包含：ETF产品规模、份额变动、净值表现等
表5通常包含：ETF产品交易数据、流动性指标等

输出要求：
- 使用中文书写
- 为表4和表5分别生成Markdown表格
- 计算并标注关键变动（如规模增减、净值涨跌等）
- 在每个表格下方附简要分析说明
- 重点突出异常变动和显著趋势
- 数据精确，保留适当小数位
- 语气正式、客观、专业
- 如果数据中有明显的趋势或异常，在分析中指出

请直接输出整理好的表4和表5内容，无需额外解释。"""

    def run_with_excel(self, excel_path: str) -> str:
        """Process ETF Excel file into Tables 4-5 for the report."""
        excel_text = parse_excel_to_text(excel_path)

        if excel_text.startswith("[ERROR]"):
            print(f"  [ETFAgent] Warning: {excel_text}")
            return excel_text

        prompt = f"""请根据以下从Excel提取的ETF数据，生成双周报的表4和表5内容：

---Excel数据开始---
{excel_text}
---Excel数据结束---

请分别输出表4和表5的完整内容（含Markdown表格和分析说明）。"""
        return self.run(prompt)

    def run_with_text(self, data_text: str) -> str:
        """Process ETF data provided as text (fallback if no Excel)."""
        prompt = f"""请根据以下ETF数据，生成双周报的表4和表5内容：

---数据开始---
{data_text}
---数据结束---

请分别输出表4和表5的完整内容（含Markdown表格和分析说明）。"""
        return self.run(prompt)
